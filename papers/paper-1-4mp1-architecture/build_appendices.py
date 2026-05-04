#!/usr/bin/env python3
"""
Build script for Paper #1 appendices G, H, I.

Generates a working docx with:
- Appendix G — Reference Implementation Naming and Mapping (concrete content)
- Appendix H — Industry Benchmarks (placeholder structure)
- Appendix I — Internal Benchmarks (placeholder structure)

Output: paper-1-appendices-G-H-I.docx (next to this script).
Edit the data dictionaries below; rerun the script to regenerate the docx.
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


HERE = Path(__file__).resolve().parent
OUTPUT_PATH = HERE / "paper-1-appendices-G-H-I.docx"


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)


def add_paragraph(doc: Document, text: str, *, italic: bool = False, bold: bool = False) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.italic = italic
    run.bold = bold
    run.font.size = Pt(11)


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run = p.add_run(text)
    run.italic = True
    run.font.size = Pt(10)


def shade_cell(cell, fill_hex: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_hex)
    tc_pr.append(shd)


def make_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = True

    header_row = table.rows[0]
    for idx, label in enumerate(headers):
        cell = header_row.cells[idx]
        cell.text = ""
        para = cell.paragraphs[0]
        run = para.add_run(label)
        run.bold = True
        run.font.size = Pt(10)
        shade_cell(cell, "D9E2F3")

    for r_idx, row in enumerate(rows, start=1):
        for c_idx, value in enumerate(row):
            cell = table.rows[r_idx].cells[c_idx]
            cell.text = ""
            para = cell.paragraphs[0]
            run = para.add_run(value)
            run.font.size = Pt(10)


# ---------------------------------------------------------------------------
# Appendix G — Reference Implementation Naming and Mapping
# ---------------------------------------------------------------------------

NAMING_MAP_INTRO = (
    "The body of this paper describes the framework, the architecture, and the "
    "verification chain in abstract terms. This appendix maps every abstract "
    "term used in the paper to the concrete name and artefact in the reference "
    "implementation maintained by the authors. Independent implementations of "
    "the 4M+1 framework are expected; the abstract names in the paper body are "
    "the canonical ones. The names in this appendix identify a specific "
    "reference implementation suitable for the case-study and benchmark "
    "appendices that follow."
)

NAMING_MAP_HEADERS = ["Paper-body abstract term", "Reference-implementation name", "Type", "Notes"]

NAMING_MAP_ROWS = [
    ["The 4M+1 framework", "4M+1", "Framework name", "Identical — 4M+1 is the contribution name, not internal-only"],
    ["The canonical entry-point file", "maestro.md", "File", "Markdown with YAML frontmatter; harness-agnostic"],
    ["The governance plane / orchestrator process", "omaestro", "Repository / process", "agent-harness-engineering/omaestro"],
    ["The security peer agent", "sec-agent", "Repository / process", "agent-harness-engineering/sec-agent — separate sidecar via gRPC over Unix socket"],
    ["The quality peer agent", "qa-agent", "Repository / process", "agent-harness-engineering/qa-agent (planned; ologos-ai's qa-agent/agent.py seeds v1)"],
    ["A predecessor 4M-tradition harness (lineage analysis)", "ThinxS", "Repository", "agent-harness-engineering/ThinxS — primary 4M development repo"],
    ["The reference deployment / operator system", "ologos-ai", "Repository / deployment", "ologos-corp/ologos-ai — reference deployment exemplar"],
    ["The Claude Code-shaped harness stub", "CLAUDE.md", "File", "One line: @maestro.md — Means-adapter file, not a framework component"],
    ["The cross-tool standard harness stub", "AGENTS.md", "File", "Brief markdown pointer to maestro.md — public AGENTS.md convention"],
    ["The five governance / Means module files", "mission.md / mind.md / morals.md / memory.md / means.md", "Files", "At repo root, peer-organized"],
    ["The Means adapter for Claude Code", "adapters/claude-code.md", "File", "Per-harness adapter; declares supports_modes + adapter identity"],
    ["The provider abstraction registry", "profiles.yaml", "File", "Sibling to maestro.md; required for model_mode types single and multi"],
    ["The multi-mode routing rules", "routers.yaml", "File", "Sibling to maestro.md; required for model_mode type multi (v1.x)"],
    ["The governance ledger (Stream A)", "audit/governance.jsonl", "Append-only file", "File-level schema-versioned; tamper-evident; run-finalized at exit"],
    ["The operational trace (Stream B)", "OTel spans via OTLP", "Wire stream", "Default exporter: Arize Phoenix self-hosted"],
    ["The wire protocol", "stdio JSON-RPC on FDs 3/4", "Protocol", "Per-agent-run lifecycle; fail-stop on disconnect or timeout"],
]

NAMING_MAP_REPO_TABLE_HEADERS = ["Repository", "URL", "Visibility", "Role"]

NAMING_MAP_REPO_TABLE_ROWS = [
    ["agent-harness-engineering/omaestro", "(redacted in public preprint; Zenodo DOI to be minted)", "Private (initial)", "Reference implementation of the governance plane"],
    ["agent-harness-engineering/sec-agent", "(redacted in public preprint; Zenodo DOI to be minted)", "Private (initial)", "Reference implementation of the security peer agent"],
    ["agent-harness-engineering/qa-agent", "(planned)", "Private (initial)", "Reference implementation of the quality peer agent"],
    ["agent-harness-engineering/ThinxS", "(redacted in public preprint)", "Private", "Predecessor 4M-tradition harness; lineage analysis"],
    ["agent-harness-engineering/harness-engineering", "(public)", "Public", "This research programme; theory, synthesis, evaluations"],
    ["ologos-corp/ologos-ai", "(redacted in public preprint)", "Private", "Reference deployment / operator system"],
]


# ---------------------------------------------------------------------------
# Appendix G.4 — Module Roster (modular federated digital enterprise framing)
# ---------------------------------------------------------------------------

MODULE_ROSTER_INTRO = (
    "The reference implementation is structured as a modular federated digital "
    "enterprise: omaestro is the governance plane for an N-module enterprise, "
    "where each capability domain attaches as an independently-governed module "
    "with its own manifest, gates, tool registry, and model-mode choice. Shared "
    "infrastructure (one orchestrator binary, one wire protocol, one verification "
    "chain, one peer-agent ecosystem) underlies all modules. The v1 reference "
    "deployment ships modules 1 and 2; subsequent modules attach without "
    "architectural change. This subsection enumerates the planned module roster."
)

MODULE_ROSTER_HEADERS = ["#", "Module", "Telos", "Likely model-mode", "Distinguishing constraint", "v1 status"]

MODULE_ROSTER_ROWS = [
    ["1", "Dev", "Code work — software engineering, code review, test authoring, build orchestration", "harness_owned (frontier reasoning)", "Tight code-quality gates; secrets and force-push emphasis", "v1.0 — reference deployment"],
    ["2", "Ops / Infra", "System administration — deployment, configuration management, infrastructure as code", "single (local LLM)", "High-volume, cost-sensitive; tight destructive-action gates", "v1.0 — reference deployment"],
    ["3", "Digital Engineering", "Systems engineering and MBSE — model-based design, requirements traceability, verification", "single or multi", "Long-running model-based artifacts; traceability across artifact lifecycle", "Future"],
    ["4", "ERP", "Business processes — finance, procurement, HR, supplier management", "single (audited cloud)", "Compliance-heavy; strict audit retention; PII gates", "Future"],
    ["5", "Mod-Sim", "Modeling and simulation — physics, finance, operations research", "single specialized", "Specialized model selection; long compute jobs; GPU/HPC profile awareness", "Future"],
    ["N+", "Future capabilities", "TBD per organization", "Per-module choice", "Same federated pattern; no architecture change required", "Future"],
]


# ---------------------------------------------------------------------------
# Appendix H — Industry Benchmarks (placeholder structure)
# ---------------------------------------------------------------------------

INDUSTRY_BENCHMARKS_INTRO = (
    "This appendix structures the comparative evaluation of the 4M+1 framework "
    "and its reference implementation against industry alternatives. v1 of this "
    "paper publishes the appendix with the structures defined and rows "
    "identified, but with measurements pending. Subsequent versioned releases "
    "(each minting a new DOI) will populate the values as data becomes "
    "available, with measurement methodology documented per row. This "
    "appendix is intentionally cross-paper — papers later in the series may "
    "extend the same tables with additional rows."
)

INDUSTRY_TABLE_1_HEADERS = ["Capability", "4M+1 (this work)", "NemoClaw / OpenClaw", "Constitutional AI", "RLHF gates", "MCP", "Notes"]

INDUSTRY_TABLE_1_ROWS = [
    ["Semantic gate enforcement", "[v1]", "[TBD]", "[TBD]", "[TBD]", "[TBD]", "BLOCK / CONFIRM / LOG taxonomy"],
    ["Structural sandbox enforcement", "[interop only]", "[TBD]", "[TBD]", "[TBD]", "[TBD]", "4M+1 composes; does not implement"],
    ["Cryptographic audit chain", "[v1]", "[TBD]", "[TBD]", "[TBD]", "[TBD]", "Five-actor Ed25519 model"],
    ["Per-call model attribution", "[v1]", "[TBD]", "[TBD]", "[TBD]", "[TBD]", "In both governance ledger and OTel trace"],
    ["Harness portability", "[v1]", "[TBD]", "[TBD]", "[TBD]", "[TBD]", "Harness-stub pattern + Means adapter"],
    ["Model-mode independence", "[v1]", "[TBD]", "[TBD]", "[TBD]", "[TBD]", "harness_owned / single / multi"],
    ["Peer-agent decomposition (DevSecOps roles)", "[v1]", "[TBD]", "[TBD]", "[TBD]", "[TBD]", "Sec, QA, future Ops"],
    ["Self-modification prohibition", "[v1]", "[TBD]", "[TBD]", "[TBD]", "[TBD]", "Explicit by absence of code"],
    ["Operator authentication standardised", "[v1]", "[TBD]", "[TBD]", "[TBD]", "[TBD]", "Pluggable IdentityProvider"],
    ["Tier 3 hardware-attestation hooks", "[v1 reserved; v2 wired]", "[TBD]", "[TBD]", "[TBD]", "[TBD]", "Schema fields present in v1"],
]

INDUSTRY_TABLE_2_HEADERS = ["Comparator", "Audit format", "Tamper-evident?", "Standardised wire?", "Independent peer agents?", "Notes"]

INDUSTRY_TABLE_2_ROWS = [
    ["4M+1 (this work)", "JSONL + OTel", "Yes (Stream A signed, append-only)", "Yes (stdio JSON-RPC)", "Yes (sec-agent + qa-agent)", "Reference implementation"],
    ["NemoClaw / OpenClaw", "[TBD]", "[TBD]", "[TBD]", "[TBD]", "Structural focus; complementary to semantic"],
    ["Constitutional AI", "[TBD]", "[TBD]", "[TBD]", "[TBD]", "Training-time, not runtime governance"],
    ["RLHF (general)", "[TBD]", "[TBD]", "[TBD]", "[TBD]", "Training-time, not auditable per-action"],
    ["MCP", "[TBD]", "[TBD]", "[TBD]", "[TBD]", "Tool protocol, not governance"],
    ["LangSmith / Braintrust / Weave / Phoenix", "[TBD]", "[TBD]", "[TBD]", "[TBD]", "Observability, not governance"],
]


# ---------------------------------------------------------------------------
# Appendix I — Internal Benchmarks (placeholder structure)
# ---------------------------------------------------------------------------

INTERNAL_BENCHMARKS_INTRO = (
    "This appendix structures the measured performance of the 4M+1 reference "
    "implementation. v1 of this paper publishes the appendix with the metrics "
    "defined and the methodology specified, but with measurements pending — "
    "the reference implementation is in active development per the phased plan "
    "(see Appendix E). Subsequent versioned releases will populate the values "
    "as the implementation phases land. Methodology notes accompany each "
    "metric so the measurement is reproducible against any conforming "
    "implementation."
)

INTERNAL_TABLE_1_HEADERS = ["Metric", "Median", "p99", "Methodology", "Phase landed", "Notes"]

INTERNAL_TABLE_1_ROWS = [
    ["Wire-protocol handshake latency (FDs 3/4 setup → first verdict ready)", "[TBD]", "[TBD]", "10k synthetic handshakes; per-run sidecar lifecycle; warm filesystem cache; stable hardware reference profile (TBD spec)", "P1", "Per-agent-run cost"],
    ["Gate evaluation latency (single gate, in-process)", "[TBD]", "[TBD]", "10k synthetic actions; mixed gate types; reference profile", "P3", "Excludes peer-agent IPC"],
    ["Peer-agent IPC overhead (gRPC over Unix socket)", "[TBD]", "[TBD]", "10k synthetic gate evaluations consulting one peer; warm connection; reference profile", "P6", "Adds to gate evaluation latency for peer-routed gates"],
    ["Audit ledger growth rate (Stream A)", "[TBD]", "[TBD]", "Per-agent-run measurements across representative workloads; bytes per gate evaluation; bytes per consequential action", "P4", "Drives storage planning"],
    ["OTel span emission overhead", "[TBD]", "[TBD]", "10k spans emitted with default Phoenix exporter; batched mode; reference profile", "P4", "Stream B sampling configurable"],
    ["Verification chain replay time (downstream reviewer)", "[TBD]", "[TBD]", "Run-finalised manifest with N=1k audit entries; signature verification per entry; cold cache", "P4", "Auditability cost"],
    ["Peer-agent crash recovery time (omaestro detection → fail-stop)", "[TBD]", "[TBD]", "Synthetic crash injection at random points during run; measure detection-to-fail-stop interval", "P6", "Critical for fail-stop semantics"],
    ["Sandbox capability probe overhead at handshake", "[TBD]", "[TBD]", "Probes for full requires_sandbox capability set; reference sandbox configurations (Landlock + seccomp; gVisor; Firecracker)", "P7", "Per-handshake one-time cost"],
    ["Manifest load time (parse + validate + verify)", "[TBD]", "[TBD]", "Cold-start manifest load with realistic five-module corpus; YAML frontmatter strict-mode parsing; signature verification", "P0 / P2", "Drives per-run startup cost"],
    ["Agent-side wrapper overhead (pkg/omaestrolib vs. direct wire)", "[TBD]", "[TBD]", "Compare typed-Go convenience client vs. raw stdio JSON-RPC under identical workloads", "P1", "Library transparency check"],
]

INTERNAL_TABLE_2_HEADERS = ["Resource", "Steady-state", "Burst", "Methodology", "Notes"]

INTERNAL_TABLE_2_ROWS = [
    ["RAM (orchestrator process, idle)", "[TBD]", "[TBD]", "Resident set size; reference profile; 60s steady-state observation; 30s burst under workload (TBD)", "Drives co-location decisions"],
    ["RAM (per peer agent, idle)", "[TBD]", "[TBD]", "Same methodology, per peer; sec-agent and qa-agent measured separately", "Independent peer crash blast"],
    ["CPU (orchestrator, gate evaluation)", "[TBD]", "[TBD]", "% of one core under realistic workload; reference profile", "Capacity planning"],
    ["Disk I/O (Stream A append rate)", "[TBD]", "[TBD]", "MB/min during representative workload; reference profile", "Storage tier selection"],
    ["Network I/O (Stream B OTLP export)", "[TBD]", "[TBD]", "MB/min to Phoenix exporter; default sampling", "Stream B can be tuned per deployment"],
]

INTERNAL_TABLE_3_HEADERS = ["Scenario", "Outcome", "Methodology", "Notes"]

INTERNAL_TABLE_3_ROWS = [
    ["sec-agent process crash mid-run", "[TBD]", "Inject crash at random gate evaluation; verify omaestro detects via gRPC connection loss; verify fail-stop", "Architectural commitment: peer crash does not crash orchestrator"],
    ["qa-agent process crash mid-run", "[TBD]", "Same methodology with qa-agent target", "Symmetry with sec-agent"],
    ["omaestro process crash mid-run", "[TBD]", "Inject crash; verify agent fail-stops on disconnect (no cached verdicts, no fallback to allow)", "Per (a) execution model"],
    ["Stream A write failure", "[TBD]", "Inject write-error on governance ledger; verify orchestrator fail-stops the agent", "Per (i): Stream A is authoritative"],
    ["Stream B export failure", "[TBD]", "Inject OTLP exporter failure; verify orchestrator continues; verify Stream A meta-event recorded", "Per (i): Stream B failure does not block agent action"],
    ["Sandbox attestation mismatch at handshake", "[TBD]", "Run with requires_sandbox declared but structural layer absent or weaker; verify fail-stop at handshake", "Per (j): fail-closed when declared-and-mismatched"],
    ["Adapter signature failure at load", "[TBD]", "Tampered adapter file; verify omaestro rejects at handshake", "Per (c)"],
    ["Operator signature failure at manifest load", "[TBD]", "Tampered manifest signature; verify omaestro rejects at handshake", "Per (c)"],
    ["Peer-agent signature failure", "[TBD]", "Mismatched key on peer-agent verdict; verify orchestrator rejects the verdict and fail-stops", "Per (c) and peer-agent ecosystem"],
    ["Wire-protocol version mismatch", "[TBD]", "Run with newer manifest schema against older orchestrator; verify fail-fast", "Per (b)"],
]


# ---------------------------------------------------------------------------
# Document assembly
# ---------------------------------------------------------------------------

def build_document() -> Document:
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)

    title = doc.add_heading("Verifiable Governance for AI Agents: The 4M+1 Framework as Enterprise Architecture", level=0)
    for run in title.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)

    p = doc.add_paragraph()
    p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run = p.add_run("Working draft — appendices G, H, I (paper 1 of series) — 2026-05-04")
    run.italic = True
    run.font.size = Pt(10)

    doc.add_paragraph()
    add_paragraph(
        doc,
        "This document carries the three concrete-content appendices for paper 1 — the "
        "naming map and the benchmark scaffolding. The body of the paper describes the "
        "framework, architecture, and verification chain in abstract terms; this "
        "appendix layer makes the abstract terms cross-referenceable to a specific "
        "reference implementation, and reserves the structure for industry- and "
        "internally-measured benchmarks that subsequent versioned releases will "
        "populate. v1 of this paper publishes with placeholders; subsequent DOIs are "
        "minted as values fill in.",
        italic=True,
    )

    doc.add_page_break()

    # Appendix G
    add_heading(doc, "Appendix G — Reference Implementation Naming and Mapping", level=1)
    add_paragraph(doc, NAMING_MAP_INTRO)
    doc.add_paragraph()

    add_heading(doc, "G.1 Term Mapping", level=2)
    make_table(doc, NAMING_MAP_HEADERS, NAMING_MAP_ROWS)
    add_caption(doc, "Table G.1 — Abstract terms used in the paper body and their reference-implementation names.")
    doc.add_paragraph()

    add_heading(doc, "G.2 Repository Inventory", level=2)
    add_paragraph(
        doc,
        "Repository URLs and DOIs are redacted in the public preprint where the "
        "underlying repository is currently private. Each will receive a Zenodo DOI "
        "minted at release; this table is the canonical reference once minted.",
    )
    doc.add_paragraph()
    make_table(doc, NAMING_MAP_REPO_TABLE_HEADERS, NAMING_MAP_REPO_TABLE_ROWS)
    add_caption(doc, "Table G.2 — Reference-implementation repository inventory.")
    doc.add_paragraph()

    add_heading(doc, "G.3 Independent Implementations", level=2)
    add_paragraph(
        doc,
        "Independent implementations of the 4M+1 framework are explicitly invited. "
        "An implementation is conformant if it (i) exposes a canonical entry-point "
        "file analogous to maestro.md and the five module files; (ii) implements the "
        "wire-protocol contract specified in Appendix B; (iii) emits a Stream A "
        "governance ledger conformant with Appendix C and a Stream B trace via OTLP; "
        "(iv) supports the harness-stub pattern such that agent tools' auto-load "
        "conventions are satisfiable by thin pointer files. Conformant "
        "implementations may use different concrete names, identity providers, or "
        "transport choices within the constraints documented above.",
    )
    doc.add_paragraph()

    add_heading(doc, "G.4 Module Roster (Modular Federated Digital Enterprise)", level=2)
    add_paragraph(doc, MODULE_ROSTER_INTRO)
    doc.add_paragraph()
    make_table(doc, MODULE_ROSTER_HEADERS, MODULE_ROSTER_ROWS)
    add_caption(doc, "Table G.4 — Reference deployment module roster. v1 deployment ships modules 1 and 2; subsequent modules attach without architectural change.")
    doc.add_paragraph()
    add_paragraph(
        doc,
        "Per-module independence and shared infrastructure are made explicit in the body of "
        "the paper (Section 6.1, Strategic Viewpoint, and Section 6.4, Resources / Services "
        "Viewpoint). The reference deployment realizing this roster is documented in the "
        "deployment-plan tracker (referenced in the repository inventory of Table G.2).",
    )

    doc.add_page_break()

    # Appendix H
    add_heading(doc, "Appendix H — Industry Benchmarks", level=1)
    add_paragraph(doc, INDUSTRY_BENCHMARKS_INTRO)
    doc.add_paragraph()

    add_heading(doc, "H.1 Capability Coverage Matrix", level=2)
    add_paragraph(
        doc,
        "Capability coverage of the 4M+1 framework alongside industry alternatives. "
        "Marked [v1] indicates a capability shipped by 4M+1's reference implementation "
        "in version 1; [TBD] indicates measurements or claims pending publication.",
    )
    doc.add_paragraph()
    make_table(doc, INDUSTRY_TABLE_1_HEADERS, INDUSTRY_TABLE_1_ROWS)
    add_caption(doc, "Table H.1 — Capability coverage matrix (placeholder; values populate in subsequent revisions).")
    doc.add_paragraph()

    add_heading(doc, "H.2 Architectural Properties Comparison", level=2)
    add_paragraph(
        doc,
        "Architectural properties relevant to verifiability and operator trust. The "
        "comparison is qualitative; each [TBD] cell anchors a measurement or "
        "documented-claim entry to be produced in subsequent paper revisions.",
    )
    doc.add_paragraph()
    make_table(doc, INDUSTRY_TABLE_2_HEADERS, INDUSTRY_TABLE_2_ROWS)
    add_caption(doc, "Table H.2 — Architectural properties comparison (placeholder).")
    doc.add_paragraph()

    add_heading(doc, "H.3 Methodology Notes", level=2)
    add_paragraph(
        doc,
        "Industry comparisons are produced from public documentation, peer-reviewed "
        "publications, and reproduction-where-possible against published interfaces. "
        "Each populated cell will reference the source of the comparison datum; for "
        "comparators that do not publish equivalent measurements, the corresponding "
        "row will be marked \"not published\" rather than left as [TBD].",
    )

    doc.add_page_break()

    # Appendix I
    add_heading(doc, "Appendix I — Internal Benchmarks", level=1)
    add_paragraph(doc, INTERNAL_BENCHMARKS_INTRO)
    doc.add_paragraph()

    add_heading(doc, "I.1 Latency Metrics", level=2)
    add_paragraph(
        doc,
        "Latency metrics for the reference implementation. Phase column references "
        "the implementation phase (Appendix E) at which the measurement becomes "
        "possible. All measurements taken on a documented reference profile (TBD: "
        "specific hardware, OS, kernel version, Go version, governance ledger "
        "filesystem class).",
    )
    doc.add_paragraph()
    make_table(doc, INTERNAL_TABLE_1_HEADERS, INTERNAL_TABLE_1_ROWS)
    add_caption(doc, "Table I.1 — Latency metrics for the reference implementation (placeholder).")
    doc.add_paragraph()

    add_heading(doc, "I.2 Resource Footprint", level=2)
    add_paragraph(
        doc,
        "Steady-state and burst resource consumption of the orchestrator and peer "
        "agents under representative workloads.",
    )
    doc.add_paragraph()
    make_table(doc, INTERNAL_TABLE_2_HEADERS, INTERNAL_TABLE_2_ROWS)
    add_caption(doc, "Table I.2 — Resource footprint of the reference implementation (placeholder).")
    doc.add_paragraph()

    add_heading(doc, "I.3 Failure-Mode Coverage", level=2)
    add_paragraph(
        doc,
        "Each architectural failure mode and its expected outcome under the locked "
        "decisions. Verified outcomes populate as implementation phases land.",
    )
    doc.add_paragraph()
    make_table(doc, INTERNAL_TABLE_3_HEADERS, INTERNAL_TABLE_3_ROWS)
    add_caption(doc, "Table I.3 — Failure-mode coverage for the reference implementation (placeholder).")
    doc.add_paragraph()

    add_heading(doc, "I.4 Methodology Notes", level=2)
    add_paragraph(
        doc,
        "All internal benchmarks are run against a documented reference profile; "
        "measurement scripts and harnesses live alongside the reference "
        "implementation and produce machine-readable output. The intent is "
        "reproducibility — a third party building a conformant implementation "
        "can run the same benchmark suite against their implementation and "
        "compare. Methodology revisions versioned alongside paper DOIs.",
    )

    return doc


def main() -> None:
    doc = build_document()
    doc.save(str(OUTPUT_PATH))
    print(f"Wrote {OUTPUT_PATH} ({OUTPUT_PATH.stat().st_size:,} bytes)")


if __name__ == "__main__":
    main()
